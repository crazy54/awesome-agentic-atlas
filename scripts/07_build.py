"""Build the themed Awesome Agent Orchestrators workbooks (dark + light).

Colors come from the validated 8-slot categorical palette (one slot per
category, fixed order, never cycled). Category identity is always carried by
the category *name* as text plus a saturated left rail -- never by fill alone --
which satisfies the relief rule for the three light-mode slots that sit under
3:1 against the light surface.
"""
import json
import math
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as TFont,
    Paragraph,
    ParagraphProperties,
    RichTextProperties,
)
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "cache"
OUT = ROOT

SOURCE = "https://github.com/andyrewlee/awesome-agent-orchestrators"
EMU = 9525  # EMU per pixel

SHOT_W, SHOT_H = 300, 150          # screenshot cell, in px
CONTENT_W, CONTENT_H = 288, 138    # image content inside it (leaves a 6px margin)

# Validated categorical palette: slot order is the CVD-safety mechanism.
CAT_ORDER = [
    ("terminal",   "Parallel Coding Agents — Terminal (TUI/CLI)", "#2a78d6", "#3987e5"),
    ("desktop",    "Parallel Coding Agents — Desktop & Web",      "#eb6834", "#d95926"),
    ("swarms",     "Multi-Agent Swarms",                          "#1baf7a", "#199e70"),
    ("loops",      "Autonomous Loop Runners",                     "#eda100", "#c98500"),
    ("tasks",      "Autonomous Task Runners",                     "#e87ba4", "#d55181"),
    ("infra",      "Agent Infrastructure & Primitives",           "#008300", "#008300"),
    ("assistants", "Personal Assistants",                         "#4a3aa7", "#9085e9"),
    ("resting",    "Resting (no recent pushes)",                  "#e34948", "#e66767"),
]
CAT_BLURB = {
    "terminal": "tmux panes, worktrees and TUI dashboards — run several agents from a terminal",
    "desktop": "the same parallel-session workflow as a desktop, browser or phone app",
    "swarms": "specialist agents that actively coordinate and delegate toward one goal",
    "loops": "drive a single goal through a retry-until-verified loop",
    "tasks": "unattended runs pulled from an issue queue, board or schedule",
    "infra": "control planes, protocols and runtimes beneath your agents",
    "assistants": "always-on agents you reach over chat; they remember across sessions",
    "resting": "no push in the last few months — parked until they wake up",
}

THEMES = {
    "dark": dict(
        surface="1A1A19", plane="0D0D0D", band="212120", header="000000",
        ink="FFFFFF", ink2="C3C2B7", muted="898781", grid="2C2C2A",
        link="86B6EF", bar="3987E5", chip_mix=0.22, chip_ink_from_hue=True,
        good="0CA30C", warn="FAB219", off="6E6D68", crit="D03B3B",
        tab="1A1A19", shot_bg=(26, 26, 25), slot=3,
    ),
    "light": dict(
        surface="FCFCFB", plane="F1F0EC", band="F6F5F1", header="12120F",
        ink="0B0B0B", ink2="52514E", muted="898781", grid="E1E0D9",
        link="1C5CAB", bar="2A78D6", chip_mix=0.86, chip_ink_from_hue=False,
        good="0A7C0A", warn="8A5A00", off="898781", crit="B02F2F",
        tab="FFFFFF", shot_bg=(252, 252, 251), slot=2,
    ),
}

COLUMNS = [
    ("#",             5.0,  "center"),
    ("Category",     26.0,  "left"),
    ("Tool",         24.0,  "left"),
    ("Screenshot",   42.14, "center"),
    ("★ Stars",      11.0,  "right"),
    ("Rank",          7.0,  "center"),
    ("What It Does", 60.0,  "left"),
    ("Win\nNative",  10.0,  "center"),
    ("Win\nWSL2",    10.0,  "center"),
    ("macOS",        10.0,  "center"),
    ("Linux",        10.0,  "center"),
    ("Docker",       10.0,  "center"),
    ("GH\nAction",   10.0,  "center"),
    ("OS Support",   26.0,  "left"),
    ("Conf.",         9.0,  "center"),
    ("Why (evidence)", 46.0, "left"),
    ("Install / Run", 54.0, "left"),
    ("Method",       11.0,  "center"),
    ("Language",     13.0,  "left"),
    ("License",      11.0,  "center"),
    ("Last Push",    12.0,  "center"),
    ("Repo",         34.0,  "left"),
    ("Website",      30.0,  "left"),
]


# ----------------------------------------------------------------- color utils
def hx(c: str) -> tuple[int, int, int]:
    c = c.lstrip("#")
    return tuple(int(c[i : i + 2], 16) for i in (0, 2, 4))


def to_hex(rgb: tuple[int, int, int]) -> str:
    return "".join(f"{max(0, min(255, int(round(v)))):02X}" for v in rgb)


def mix(a: str, b: str, t: float) -> str:
    """t=0 -> a, t=1 -> b."""
    ra, rb = hx(a), hx(b)
    return to_hex(tuple(ra[i] + (rb[i] - ra[i]) * t for i in range(3)))


def _lin(v: float) -> float:
    v /= 255
    return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4


def luminance(c: str) -> float:
    r, g, b = hx(c)
    return 0.2126 * _lin(r) + 0.7152 * _lin(g) + 0.0722 * _lin(b)


def contrast(a: str, b: str) -> float:
    la, lb = luminance(a), luminance(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def best_ink(bg: str, inks: list[str]) -> tuple[str, float]:
    scored = sorted(((contrast(bg, i), i) for i in inks), reverse=True)
    return scored[0][1], scored[0][0]


# ----------------------------------------------------------------- helpers
def fill(c: str) -> PatternFill:
    return PatternFill("solid", start_color=c, end_color=c)


def side(c: str, style: str = "thin") -> Side:
    return Side(style=style, color=c)


def pad_shot(src: Path, bg: tuple[int, int, int]) -> BytesIO:
    """Center the thumbnail on an exact SHOT_W x SHOT_H canvas in the theme bg.

    Uniform image size means a twoCell anchor can stretch each picture to the
    same cell without distortion -- and twoCell is what makes pictures hide
    along with their row when the user applies a filter.
    """
    canvas = PILImage.new("RGB", (SHOT_W, SHOT_H), bg)
    try:
        im = PILImage.open(src)
        im.load()
        im = im.convert("RGB")
        im.thumbnail((CONTENT_W, CONTENT_H), PILImage.LANCZOS)
        canvas.paste(im, ((SHOT_W - im.width) // 2, (SHOT_H - im.height) // 2))
    except Exception:
        pass
    buf = BytesIO()
    canvas.save(buf, "JPEG", quality=88, optimize=True)
    buf.seek(0)
    return buf


def paint(ws, r1, r2, c1, c2, style):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).fill = style


# ----------------------------------------------------------------- main sheet
def build_main(wb, recs, T, cat_style, shots, pool=None):
    ws = wb.create_sheet("Orchestrators")
    ws.sheet_properties.tabColor = T["surface"]
    ws.sheet_view.showGridLines = False

    ncols = len(COLUMNS)
    last_row = 2 + len(recs)

    # Theme the visible canvas, including a margin past the data.
    paint(ws, 1, last_row + 40, 1, ncols + 6, fill(T["surface"]))

    for i, (name, width, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # ---- row 1: title band
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws.merge_cells(start_row=1, start_column=8, end_row=1, end_column=ncols)
    t = ws.cell(row=1, column=1, value="AWESOME AGENT ORCHESTRATORS")
    t.font = Font(name="Segoe UI", size=15, bold=True, color=T["ink"])
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    m = ws.cell(row=1, column=8,
                value=f"{len(recs)} tools  ·  grouped by category, then most→fewest stars  ·  "
                      f"snapshot {date.today().isoformat()}   ")
    m.font = Font(name="Segoe UI", size=10, color=T["muted"])
    m.alignment = Alignment(horizontal="right", vertical="center")
    paint(ws, 1, 1, 1, ncols, fill(T["plane"]))
    ws.row_dimensions[1].height = 30

    # ---- row 2: header
    for i, (name, _, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=2, column=i, value=name)
        c.fill = fill(T["header"])
        c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=side(T["bar"], "medium"))
    ws.row_dimensions[2].height = 32

    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{last_row}"

    # ---- data rows
    f_body = Font(name="Segoe UI", size=10, color=T["ink"])
    f_body2 = Font(name="Segoe UI", size=10, color=T["ink2"])
    f_muted = Font(name="Segoe UI", size=9, color=T["muted"])
    f_mono = Font(name="Consolas", size=9, color=T["ink"])
    f_tool = Font(name="Segoe UI", size=11, bold=True, color=T["link"], underline="single")
    f_link = Font(name="Segoe UI", size=9, color=T["link"], underline="single")
    f_stars = Font(name="Segoe UI", size=12, bold=True, color=T["ink"])

    status_font = {
        "Yes": Font(name="Segoe UI", size=10, bold=True, color=T["good"]),
        "Likely": Font(name="Segoe UI", size=10, color=T["warn"]),
        "No": Font(name="Segoe UI", size=10, color=T["off"]),
        "n/a": Font(name="Segoe UI", size=10, color=T["muted"]),
    }
    hair = side(T["grid"])

    for idx, r in enumerate(recs):
        row = 3 + idx
        ws.row_dimensions[row].height = SHOT_H * 0.75  # px -> pt
        cs = cat_style[r["cat_key"]]
        band = T["band"] if idx % 2 else T["surface"]

        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = fill(band)
            cell.border = Border(bottom=hair)

        top = Alignment(vertical="top", wrap_text=True)
        ctr = Alignment(horizontal="center", vertical="center")

        # A: rank within category
        a = ws.cell(row=row, column=1, value=r["cat_rank"])
        a.font = f_muted
        a.alignment = ctr
        a.border = Border(bottom=hair, left=side(cs["hue"], "thick"))

        # B: category chip -- name as text (never color alone) + saturated rail
        b = ws.cell(row=row, column=2, value=r["category_short"])
        b.fill = fill(cs["chip_bg"])
        b.font = Font(name="Segoe UI", size=9, bold=True, color=cs["chip_ink"])
        b.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        b.border = Border(bottom=hair, left=side(cs["hue"], "thick"))

        # C: tool name -> repo
        c3 = ws.cell(row=row, column=3, value=r["name"])
        c3.font = f_tool
        c3.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        c3.hyperlink = r["url"]
        if r.get("archived"):
            ws.cell(row=row, column=3).comment = None

        # D: screenshot (image inserted after styling)
        ws.cell(row=row, column=4).alignment = ctr

        # E: stars
        e = ws.cell(row=row, column=5, value=r["stars"])
        e.font = f_stars
        e.number_format = "#,##0"
        e.alignment = Alignment(horizontal="right", vertical="center", indent=1)

        f = ws.cell(row=row, column=6, value=r["overall_rank"])
        f.font = f_muted
        f.alignment = ctr

        g = ws.cell(row=row, column=7, value=r["description"])
        g.font = f_body
        g.alignment = top

        for off, key in enumerate(["win_native", "win_wsl2", "macos", "linux", "docker", "gh_action"]):
            cell = ws.cell(row=row, column=8 + off, value=r[key])
            cell.font = status_font.get(r[key], f_body)
            cell.alignment = ctr

        n = ws.cell(row=row, column=14, value=r["os_summary"])
        n.font = f_body2
        n.alignment = Alignment(vertical="center", wrap_text=True, indent=1)

        o = ws.cell(row=row, column=15, value=r["os_confidence"])
        o.font = status_font["Yes"] if r["os_confidence"] == "High" else (
            status_font["Likely"] if r["os_confidence"] == "Medium" else status_font["No"])
        o.alignment = ctr

        p = ws.cell(row=row, column=16, value=r["os_evidence"])
        p.font = f_muted
        p.alignment = top

        q = ws.cell(row=row, column=17, value=r["install_cmd"])
        q.font = f_mono
        q.alignment = Alignment(vertical="center", wrap_text=True, indent=1)

        for col, val, fnt in (
            (18, r["install_method"], f_muted),
            (19, r["language"], f_body2),
            (20, r["license"], f_muted),
            (21, r["pushed_at"], f_muted),
        ):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = fnt
            cell.alignment = ctr

        v = ws.cell(row=row, column=22, value=r["url"].replace("https://github.com/", ""))
        v.font = f_link
        v.alignment = Alignment(vertical="center", indent=1)
        v.hyperlink = r["url"]

        site = (r.get("homepage") or "").strip()
        w = ws.cell(row=row, column=23, value=site.replace("https://", "").replace("http://", "") if site else "—")
        w.alignment = Alignment(vertical="center", indent=1)
        if site:
            w.font = f_link
            w.hyperlink = site if site.startswith("http") else "https://" + site
        else:
            w.font = f_muted

    # ---- screenshots: twoCell anchor so they hide with a filtered row
    for idx, r in enumerate(recs):
        row = 3 + idx
        src = shots.get(r["nwo"], {}).get(f"shot_{'dark' if T is THEMES['dark'] else 'light'}")
        if not src or not Path(src).exists():
            continue
        anchor = TwoCellAnchor(
            editAs="twoCell",
            _from=AnchorMarker(col=3, colOff=0, row=row - 1, rowOff=0),
            to=AnchorMarker(col=4, colOff=0, row=row, rowOff=0),
        )
        # In the merged Atlas these same repos also appear on the Leaderboard, By Category and the
        # platform sheets, so `16_build_all` hands in a pool and the picture is embedded once for all
        # of them (see `media.py`). Standalone -- this script's own `build()` -- there is no other sheet
        # to share with, so the plain per-placement embed is exactly right and stays the default.
        if pool is not None:
            pool.place(ws, Path(src), anchor)
        else:
            img = XLImage(pad_shot(Path(src), T["shot_bg"]))
            img.anchor = anchor
            ws.add_image(img)

    # ---- data bar for stars (magnitude -> the sequential blue, not category hue)
    ws.conditional_formatting.add(
        f"E3:E{last_row}",
        DataBarRule(start_type="num", start_value=0, end_type="percentile", end_value=90,
                    color=T["bar"], showValue=True),
    )
    return ws


# ----------------------------------------------------------------- cover sheet
def build_cover(wb, recs, T, cat_style, stats):
    ws = wb.create_sheet("Start Here", 0)
    ws.sheet_properties.tabColor = T["bar"]
    ws.sheet_view.showGridLines = False
    paint(ws, 1, 90, 1, 14, fill(T["surface"]))

    # Tiles occupy column pairs (2,3) (4,5) ... (12,13); every pair sums to TILE_W
    # so all six tiles are the same size and the last one is not clipped.
    TILE_W = 32
    widths = [3, 22, 10, 17, 15, 17, 15, 16, 16, 16, 16, 16, 16, 6]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    assert all(widths[c] + widths[c + 1] == TILE_W for c in range(1, 13, 2)), widths
    BODY = 12  # bullets span columns 2..13, the full tile strip

    def put(row, col, val, size=10, bold=False, color=None, span=1, align="left", wrap=False, italic=False, mono=False):
        if span > 1:
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
        c = ws.cell(row=row, column=col, value=val)
        c.font = Font(name="Consolas" if mono else "Segoe UI", size=size, bold=bold,
                      italic=italic, color=color or T["ink"])
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap, indent=1 if align == "left" else 0)
        return c

    paint(ws, 1, 4, 1, 14, fill(T["plane"]))
    ws.row_dimensions[2].height = 40
    ws.row_dimensions[3].height = 20
    put(2, 2, "AWESOME AGENT ORCHESTRATORS", size=24, bold=True, span=8)
    put(3, 2, "Every tool from andyrewlee/awesome-agent-orchestrators — grouped by category, "
              "ranked by GitHub stars, with OS support, install commands and screenshots.",
        size=10, color=T["ink2"], span=9)
    put(4, 2, f"Snapshot taken {date.today().isoformat()}   ·   source list:", size=9, color=T["muted"], span=3)
    c = put(4, 5, SOURCE, size=9, color=T["link"], span=4)
    c.font = Font(name="Segoe UI", size=9, color=T["link"], underline="single")
    c.hyperlink = SOURCE

    # ---- stat tiles
    row = 6
    put(row, 2, "AT A GLANCE", size=10, bold=True, color=T["muted"])
    # The theme switch: a link to the sibling file, which lives in the same folder.
    other = "LIGHT" if T is THEMES["dark"] else "DARK"
    sw = put(row, 11, f"◐  Switch to the {other} theme  →", size=10, bold=True,
             color=T["link"], span=3, align="right")
    sw.font = Font(name="Segoe UI", size=10, bold=True, color=T["link"], underline="single")
    sw.hyperlink = f"Awesome-Agent-Orchestrators-{other}.xlsx"
    row += 1
    tiles = [
        (f"{stats['tools']}", "tools catalogued"),
        (f"{stats['stars']:,}", "GitHub stars, combined"),
        (f"{stats['cats']}", "categories"),
        (f"{stats['win']}", "run natively on Windows"),
        (f"{stats['wsl']}", "reachable via WSL2"),
        (f"{stats['shots']}", "with a screenshot"),
    ]
    for i, (big, label) in enumerate(tiles):
        col = 2 + i * 2
        last_tile = i == len(tiles) - 1
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        v = ws.cell(row=row, column=col, value=big)
        v.font = Font(name="Segoe UI", size=22, bold=True, color=T["ink"])
        v.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        l = ws.cell(row=row + 1, column=col, value=label)
        l.font = Font(name="Segoe UI", size=9, color=T["muted"])
        l.alignment = Alignment(horizontal="left", vertical="top", indent=1)
        for rr in (row, row + 1):
            for cc in (col, col + 1):
                c = ws.cell(row=rr, column=cc)
                c.fill = fill(T["band"])
                # Accent rule runs the whole tile; a surface-coloured right edge
                # is the gap that keeps neighbouring tiles from fusing.
                c.border = Border(
                    top=side(T["bar"], "medium") if rr == row else None,
                    right=side(T["surface"], "medium") if cc == col + 1 and not last_tile else None,
                )
    ws.row_dimensions[row].height = 34
    ws.row_dimensions[row + 1].height = 18

    # ---- category legend
    row += 3
    put(row, 2, "CATEGORY COLOUR KEY", size=10, bold=True, color=T["muted"])
    row += 1
    for hdr, col, span in (("Category", 2, 1), ("Tools", 3, 1), ("Stars", 4, 1), ("Top tool", 5, 2), ("What it's for", 7, 7)):
        h = put(row, col, hdr, size=9, bold=True, color="FFFFFF", span=span)
        h.fill = fill(T["header"])
        for k in range(col, col + span):
            ws.cell(row=row, column=k).fill = fill(T["header"])
    row += 1
    for key, full, _, _ in CAT_ORDER:
        cs = cat_style[key]
        s = stats["by_cat"][key]
        chip = put(row, 2, s["short"], size=9, bold=True, color=cs["chip_ink"])
        chip.fill = fill(cs["chip_bg"])
        chip.border = Border(left=side(cs["hue"], "thick"), bottom=side(T["grid"]))
        put(row, 3, s["n"], size=10, align="center")
        v = put(row, 4, s["stars"], size=10, align="right")
        v.number_format = "#,##0"
        put(row, 5, s["top"], size=9, color=T["ink2"], span=2)
        put(row, 7, CAT_BLURB[key], size=9, color=T["muted"], span=7, wrap=True)
        for cc in range(3, 14):
            ws.cell(row=row, column=cc).border = Border(bottom=side(T["grid"]))
        ws.row_dimensions[row].height = 22
        row += 1

    # Wrapped merged cells do not auto-fit in Excel, so height has to be computed.
    def bullet_height(text: str, size: int) -> float:
        span_units = sum(widths[1:1 + BODY])
        per_line = span_units * (1.30 if size <= 9 else 1.12)
        lines = max(1, math.ceil(len(text) / per_line))
        return lines * (13 if size <= 9 else 15) + 3

    # ---- how to use
    row += 1
    put(row, 2, "HOW TO FILTER", size=10, bold=True, color=T["muted"])
    row += 1
    for line in [
        'Click any dropdown arrow in the header row of the "Orchestrators" sheet.',
        'Windows-only view: filter "Win Native" → tick Yes (add Likely to widen the net).',
        'WSL2 view: filter "Win WSL2" → tick Yes. Docker, macOS, Linux and GH Action work the same way.',
        'Combine filters — e.g. Win Native = Yes plus Category = Multi-Agent Swarms.',
        'The header row and the first three columns stay locked while you scroll.',
        'Clear everything with Data ▸ Clear in the ribbon.',
    ]:
        text = "•  " + line
        put(row, 2, text, size=10, color=T["ink2"], span=BODY, wrap=True)
        ws.row_dimensions[row].height = bullet_height(text, 10)
        row += 1

    # ---- methodology
    row += 1
    put(row, 2, "HOW THE OS COLUMNS WERE DECIDED  (and where to distrust them)", size=10, bold=True, color=T["muted"])
    row += 1
    for line in [
        'Yes / Likely / No is derived evidence, not a vendor matrix. The "Why (evidence)" column shows the reasoning for every row.',
        'Strongest signal: published release assets. A .msi or windows-amd64.zip proves native Windows; a lone .dmg proves it does not.',
        'Then explicit README prose ("Windows, macOS, Linux", "macOS only", "requires WSL").',
        'Then the install route — winget/scoop imply Windows; brew implies macOS; curl|sh implies Unix, so Windows means WSL2.',
        'A hard tmux dependency sets Win Native = No and WSL2 = Yes: tmux has no native Windows build.',
        '"Likely" means a portable runtime (Node/Python/Rust/Go) with no explicit platform claim — probable, unverified.',
        'Conf. = High when assets or prose decided it, Medium for install-route inference, Low for language inference alone.',
        'GitHub Actions show n/a across the OS columns: they run on a hosted runner, so you drive them from any OS.',
        'Screenshots are the first usable image in each README; where there was none, the project website was captured live, '
        'and failing that a GitHub repo card stands in. Check the Screenshot source before trusting a card as a UI preview.',
        f'{stats["unavailable"]} repos in the source list now return 404 (deleted, renamed or made private) and are kept with zero stars so nothing silently vanishes.',
    ]:
        text = "•  " + line
        put(row, 2, text, size=9, color=T["muted"], span=BODY, wrap=True)
        ws.row_dimensions[row].height = bullet_height(text, 9)
        row += 1

    row += 1
    put(row, 2, f"Theme: {'DARK' if T is THEMES['dark'] else 'LIGHT'}. The {other} theme is the sibling "
                f"file in this folder — use the switch at the top right. (Excel cannot repaint a sheet "
                f"without macros, so the two themes are two files.)",
        size=9, italic=True, color=T["muted"], span=BODY)
    return ws


# ----------------------------------------------------------------- stats sheet
def chart_text(color: str, size: int = 900, bold: bool = False) -> RichText:
    """Chart text wears the theme's ink tokens -- never a series colour."""
    cp = CharacterProperties(solidFill=color.lstrip("#"), sz=size, b=bold,
                             latin=TFont(typeface="Segoe UI"))
    return RichText(bodyPr=RichTextProperties(),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def themed_bar(ws, T, cat_style, title: str, data: Reference, cats: Reference,
               vmax: float) -> BarChart:
    """Horizontal bar chart repainted for the sheet's theme.

    openpyxl emits an unstyled chart, which Excel renders with a white plot area and
    black gridlines -- unreadable on the dark sheet. Every surface, line and run of
    text therefore has to be set explicitly. Value axis is dropped in favour of
    direct labels (fewer marks, no axis to read against).
    """
    ch = BarChart()
    ch.type = "bar"
    ch.style = None
    ch.roundedCorners = False
    ch.varyColors = False
    ch.gapWidth = 55
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None

    ch.graphical_properties = GraphicalProperties(solidFill=T["plane"])
    ch.graphical_properties.ln = LineProperties(noFill=True)
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True)
    ch.plot_area.graphicalProperties.ln = LineProperties(noFill=True)

    ch.title = title
    tp = ch.title.tx.rich.p[0]
    cp = CharacterProperties(solidFill=T["ink"].lstrip("#"), sz=1100, b=True,
                             latin=TFont(typeface="Segoe UI"))
    tp.pPr = ParagraphProperties(defRPr=cp)
    for run in tp.r or []:
        run.rPr = cp

    # Category axis: labels only -- no ticks, hairline baseline, top-down order so
    # it reads in the same order as the table above.
    ch.x_axis.delete = False
    ch.x_axis.title = None
    ch.x_axis.majorGridlines = None
    ch.x_axis.majorTickMark = "none"
    ch.x_axis.minorTickMark = "none"
    ch.x_axis.txPr = chart_text(T["ink2"], 900)
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=T["grid"].lstrip("#"), w=9525))
    ch.x_axis.scaling.orientation = "maxMin"

    # Value axis off; the numbers sit on the bars instead. The headroom is what
    # keeps the longest label from being clipped off the end of the longest bar.
    ch.y_axis.delete = True
    ch.y_axis.title = None
    ch.y_axis.majorGridlines = None
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = vmax * 1.18

    ch.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                             showLegendKey=False, showPercent=False, showBubbleSize=False)
    ch.dLbls.numFmt = "#,##0"
    ch.dLbls.dLblPos = "outEnd"
    ch.dLbls.txPr = chart_text(T["ink"], 900, bold=True)

    ser = ch.series[0]
    ser.graphicalProperties = GraphicalProperties()
    ser.graphicalProperties.ln = LineProperties(noFill=True)
    for i, (key, *_rest) in enumerate(CAT_ORDER):
        dp = DataPoint(idx=i)
        dp.graphicalProperties.solidFill = cat_style[key]["hue"].lstrip("#")
        dp.graphicalProperties.line.noFill = True
        ser.data_points.append(dp)

    ch.width, ch.height = 13.4, 9.6
    return ch


def build_stats(wb, recs, T, cat_style, stats):
    ws = wb.create_sheet("Category Stats")
    ws.sheet_properties.tabColor = T["surface"]
    ws.sheet_view.showGridLines = False
    paint(ws, 1, 40, 1, 13, fill(T["surface"]))
    for i, w in enumerate([3, 30, 9, 12, 11, 11, 10, 10, 10, 10, 10, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("B2:L2")
    t = ws.cell(row=2, column=2, value="CATEGORY BREAKDOWN")
    t.font = Font(name="Segoe UI", size=16, bold=True, color=T["ink"])
    t.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 30
    paint(ws, 2, 2, 1, 13, fill(T["plane"]))

    hdrs = ["Category", "Tools", "Total ★", "Median ★", "Top ★", "Win", "WSL2", "macOS", "Linux", "Docker", "Top tool"]
    for i, h in enumerate(hdrs, start=2):
        c = ws.cell(row=4, column=i, value=h)
        c.fill = fill(T["header"])
        c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=side(T["bar"], "medium"))
    ws.row_dimensions[4].height = 26

    row = 5
    for key, full, _, _ in CAT_ORDER:
        s = stats["by_cat"][key]
        cs = cat_style[key]
        chip = ws.cell(row=row, column=2, value=s["short"])
        chip.fill = fill(cs["chip_bg"])
        chip.font = Font(name="Segoe UI", size=9, bold=True, color=cs["chip_ink"])
        chip.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        chip.border = Border(left=side(cs["hue"], "thick"), bottom=side(T["grid"]))
        vals = [s["n"], s["stars"], s["median"], s["max"], s["win"], s["wsl"], s["mac"], s["lin"], s["doc"]]
        for j, v in enumerate(vals, start=3):
            c = ws.cell(row=row, column=j, value=v)
            c.font = Font(name="Segoe UI", size=10, color=T["ink"])
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = Border(bottom=side(T["grid"]))
        c = ws.cell(row=row, column=12, value=s["top"])
        c.font = Font(name="Segoe UI", size=9, color=T["ink2"])
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        c.border = Border(bottom=side(T["grid"]))
        ws.row_dimensions[row].height = 22
        row += 1

    # Two single-series bar charts. Colour follows the category (the entity), never
    # its rank, so the bars match the chips in the table above them.
    cats = Reference(ws, min_col=2, min_row=5, max_row=row - 1)
    for col, title, anchor, vmax in (
        (3, "Tools per category", f"B{row + 2}", max(s["n"] for s in stats["by_cat"].values())),
        (4, "Combined GitHub stars per category", f"H{row + 2}",
         max(s["stars"] for s in stats["by_cat"].values())),
    ):
        ch = themed_bar(ws, T, cat_style, title,
                        Reference(ws, min_col=col, min_row=4, max_row=row - 1), cats, vmax)
        ws.add_chart(ch, anchor)
    return ws


# ----------------------------------------------------------------- windows sheet
def build_windows(wb, recs, T, cat_style):
    picks = [r for r in recs if r["win_native"] in ("Yes", "Likely") or r["win_wsl2"] == "Yes"]
    picks.sort(key=lambda r: (-{"Yes": 2, "Likely": 1}.get(r["win_native"], 0), -r["stars"]))

    ws = wb.create_sheet("Windows Picks")
    ws.sheet_properties.tabColor = T["bar"]
    ws.sheet_view.showGridLines = False
    last = 3 + len(picks)
    paint(ws, 1, last + 20, 1, 10, fill(T["surface"]))
    for i, w in enumerate([5, 24, 24, 11, 11, 11, 56, 54, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("A1:I1")
    t = ws.cell(row=1, column=1, value="WINDOWS PICKS  ·  runs natively on Windows, or reachable through WSL2")
    t.font = Font(name="Segoe UI", size=14, bold=True, color=T["ink"])
    t.alignment = Alignment(vertical="center", indent=1)
    paint(ws, 1, 1, 1, 9, fill(T["plane"]))
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:I2")
    s = ws.cell(row=2, column=1, value=f"{len(picks)} of {len(recs)} tools. Native first, then WSL2 — each block by stars. "
                                       "Screenshots live on the Orchestrators sheet.")
    s.font = Font(name="Segoe UI", size=9, color=T["muted"])
    s.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 18

    hdrs = ["#", "Tool", "Category", "★ Stars", "Win Native", "Win WSL2", "What It Does", "Install / Run", "Language"]
    for i, h in enumerate(hdrs, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.fill = fill(T["header"])
        c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=side(T["bar"], "medium"))
    ws.row_dimensions[3].height = 26
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:I{last}"

    status_font = {
        "Yes": Font(name="Segoe UI", size=10, bold=True, color=T["good"]),
        "Likely": Font(name="Segoe UI", size=10, color=T["warn"]),
        "No": Font(name="Segoe UI", size=10, color=T["off"]),
        "n/a": Font(name="Segoe UI", size=10, color=T["muted"]),
    }
    hair = side(T["grid"])
    for i, r in enumerate(picks):
        row = 4 + i
        cs = cat_style[r["cat_key"]]
        band = T["band"] if i % 2 else T["surface"]
        for c in range(1, 10):
            cell = ws.cell(row=row, column=c)
            cell.fill = fill(band)
            cell.border = Border(bottom=hair)
        a = ws.cell(row=row, column=1, value=i + 1)
        a.font = Font(name="Segoe UI", size=9, color=T["muted"])
        a.alignment = Alignment(horizontal="center", vertical="center")
        a.border = Border(bottom=hair, left=side(cs["hue"], "thick"))
        n = ws.cell(row=row, column=2, value=r["name"])
        n.font = Font(name="Segoe UI", size=10, bold=True, color=T["link"], underline="single")
        n.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        n.hyperlink = r["url"]
        cc = ws.cell(row=row, column=3, value=r["category_short"])
        cc.font = Font(name="Segoe UI", size=9, color=cs["chip_ink"])
        cc.fill = fill(cs["chip_bg"])
        cc.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        cc.border = Border(bottom=hair, left=side(cs["hue"], "thick"))
        st = ws.cell(row=row, column=4, value=r["stars"])
        st.font = Font(name="Segoe UI", size=11, bold=True, color=T["ink"])
        st.number_format = "#,##0"
        st.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        for off, key in enumerate(("win_native", "win_wsl2")):
            c2 = ws.cell(row=row, column=5 + off, value=r[key])
            c2.font = status_font.get(r[key], status_font["No"])
            c2.alignment = Alignment(horizontal="center", vertical="center")
        d = ws.cell(row=row, column=7, value=r["description"])
        d.font = Font(name="Segoe UI", size=9, color=T["ink"])
        d.alignment = Alignment(vertical="top", wrap_text=True)
        q = ws.cell(row=row, column=8, value=r["install_win"])
        q.font = Font(name="Consolas", size=9, color=T["ink"])
        q.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        lg = ws.cell(row=row, column=9, value=r["language"])
        lg.font = Font(name="Segoe UI", size=9, color=T["ink2"])
        lg.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 42

    ws.conditional_formatting.add(
        f"D4:D{last}",
        DataBarRule(start_type="num", start_value=0, end_type="percentile", end_value=90,
                    color=T["bar"], showValue=True),
    )
    return ws


# ----------------------------------------------------------------- assembly
def short_name(full: str) -> str:
    return {
        "Parallel Coding Agents — Terminal (TUI/CLI)": "Parallel · Terminal",
        "Parallel Coding Agents — Desktop & Web": "Parallel · Desktop & Web",
        "Multi-Agent Swarms": "Multi-Agent Swarms",
        "Autonomous Loop Runners": "Loop Runners",
        "Autonomous Task Runners": "Task Runners",
        "Agent Infrastructure & Primitives": "Infra & Primitives",
        "Personal Assistants": "Personal Assistants",
        "Resting": "Resting",
    }[full]


def prepare(recs):
    order = {k: i for i, (k, *_r) in enumerate(CAT_ORDER)}
    for r in recs:
        r["category_short"] = short_name(r["category"])
        # GitHub reports an unrecognised LICENSE file as "NOASSERTION" -- noise.
        if r["license"] in ("NOASSERTION", "NONE", "", None):
            r["license"] = "—"
        # On the Windows sheet a Homebrew line is useless. If the project ships a
        # Windows build, point at the download instead of a command that can't run.
        r["install_win"] = r["install_cmd"]
        if r["install_method"] == "brew" and r["win_native"] == "Yes":
            r["install_win"] = f"{r['url']}/releases/latest"
    recs.sort(key=lambda r: (order[r["cat_key"]], -r["stars"], r["name"].lower()))

    for rank, r in enumerate(sorted(recs, key=lambda x: -x["stars"]), start=1):
        r["overall_rank"] = rank
    seen = {}
    for r in recs:
        seen[r["cat_key"]] = seen.get(r["cat_key"], 0) + 1
        r["cat_rank"] = seen[r["cat_key"]]
    return recs


def make_stats(recs, shots):
    by_cat = {}
    for key, full, *_ in CAT_ORDER:
        rs = [r for r in recs if r["cat_key"] == key]
        st = sorted((r["stars"] for r in rs), reverse=True)
        top = max(rs, key=lambda r: r["stars"]) if rs else None
        by_cat[key] = dict(
            short=short_name(full if full in ("Multi-Agent Swarms", "Personal Assistants") else
                             next(f for _k, f, *_x in CAT_ORDER if _k == key).replace(" (no recent pushes)", "")),
            n=len(rs), stars=sum(st), median=st[len(st) // 2] if st else 0, max=st[0] if st else 0,
            top=f"{top['name']} ({top['stars']:,}★)" if top else "—",
            win=sum(1 for r in rs if r["win_native"] == "Yes"),
            wsl=sum(1 for r in rs if r["win_wsl2"] == "Yes"),
            mac=sum(1 for r in rs if r["macos"] == "Yes"),
            lin=sum(1 for r in rs if r["linux"] == "Yes"),
            doc=sum(1 for r in rs if r["docker"] == "Yes"),
        )
    return dict(
        tools=len(recs), stars=sum(r["stars"] for r in recs), cats=len(CAT_ORDER),
        win=sum(1 for r in recs if r["win_native"] == "Yes"),
        wsl=sum(1 for r in recs if r["win_wsl2"] == "Yes"),
        shots=sum(1 for r in recs if shots.get(r["nwo"], {}).get("shot_light")),
        unavailable=sum(1 for r in recs if r.get("unavailable")),
        by_cat=by_cat,
    )


def cat_styles(T):
    out, report = {}, []
    for key, full, light_hue, dark_hue in CAT_ORDER:
        hue = (dark_hue if T is THEMES["dark"] else light_hue).lstrip("#")
        if T is THEMES["dark"]:
            chip_bg = mix(T["surface"], hue, T["chip_mix"])       # hue washed toward surface
            chip_ink, ratio = best_ink(chip_bg, [hue, T["ink"], T["ink2"]])
        else:
            chip_bg = mix(hue, T["surface"], T["chip_mix"])       # pale tint of the hue
            chip_ink, ratio = best_ink(chip_bg, [T["ink"], "FFFFFF"])
        out[key] = dict(hue=hue, chip_bg=chip_bg, chip_ink=chip_ink)
        report.append((short_name(full.replace(" (no recent pushes)", "")), hue, chip_bg, chip_ink, ratio))
    return out, report


def build(theme_name, recs, shots):
    T = THEMES[theme_name]
    cat_style, report = cat_styles(T)
    stats = make_stats(recs, shots)

    wb = Workbook()
    wb.remove(wb.active)
    build_main(wb, recs, T, cat_style, shots)
    build_stats(wb, recs, T, cat_style, stats)
    build_windows(wb, recs, T, cat_style)
    build_cover(wb, recs, T, cat_style, stats)
    wb.active = 0

    path = OUT / f"Awesome-Agent-Orchestrators-{theme_name.upper()}.xlsx"
    wb.save(path)
    return path, report


def main() -> None:
    recs = prepare(json.loads((CACHE / "records.json").read_text(encoding="utf-8")))
    shots = json.loads((CACHE / "shots.json").read_text(encoding="utf-8"))

    for theme in ("dark", "light"):
        path, report = build(theme, recs, shots)
        mb = path.stat().st_size / 1e6
        print(f"\n{theme.upper()}  ->  {path.name}  ({mb:.1f} MB)")
        print(f"  {'category':28} {'hue':>8} {'chip bg':>8} {'ink':>8}  contrast")
        worst = 99.0
        for name, hue, bg, ink, ratio in report:
            flag = "OK " if ratio >= 4.5 else "LOW"
            worst = min(worst, ratio)
            print(f"  {name:28} #{hue:>7} #{bg:>7} #{ink:>7}  {ratio:5.2f}  {flag}")
        print(f"  worst chip text contrast: {worst:.2f}:1 (need >= 4.5 for body text)")


if __name__ == "__main__":
    main()
