"""Unit tests for `scripts/media.py`, and for the ZIP-entry arithmetic the fix was sized against.

The real workbook cannot be built here: stages 14+ need `cache/meta.json`, `cache/records_all.json`,
`cache/shots_all.json` and `cache/shots.json`, which come out of an authenticated GitHub crawl and are
not committed. So this file does the two things that do not need the crawl.

First it builds a synthetic workbook -- a handful of sheets showing the same few pictures, which is the
real workbook's shape in miniature -- and checks the claims the fix rests on against the archive itself
rather than against a reading of openpyxl: that a plain `wb.save` writes one media part per *placement*,
that the pool collapses those to one per distinct picture, that nothing else in the package changes, and
that the result still reads back with every drawing relationship resolving to a part that is present.

Then it writes down the projection. The ticket's premise -- that ingesting a 7,835-row list would put
the workbook past the 65,535-entry ceiling -- is arithmetic on numbers that are published, so it can be
re-run instead of re-argued, and the assumption it turns on can be varied to see which way it falls.

Run: python tests/media_test.py
"""
from __future__ import annotations

import importlib.util
import io
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from PIL import Image as PILImage

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import media  # noqa: E402

ok = bad = 0
TMP = Path(tempfile.mkdtemp(prefix="media_test_"))


def eq(name: str, got, want) -> None:
    global ok, bad
    if got == want:
        ok += 1
    else:
        bad += 1
        print(f"FAIL {name}\n  got  {got!r}\n  want {want!r}")


def true(name: str, cond, detail: str = "") -> None:
    global ok, bad
    if cond:
        ok += 1
    else:
        bad += 1
        print(f"FAIL {name}" + (f"\n  {detail}" if detail else ""))


def raises(name: str, fn) -> None:
    """SystemExit, and the message it carries -- both are the point of the guards."""
    global ok, bad
    try:
        fn()
    except SystemExit as e:
        ok += 1
        return str(e)
    bad += 1
    print(f"FAIL {name} did not raise")
    return ""


# ---- the synthetic workbook: SHEETS sheets, each showing all PICTURES of the pictures
SHEETS, PICTURES = 5, 4
PLACEMENTS = SHEETS * PICTURES

shots = TMP / "shots"
shots.mkdir()
SRC = []
for i in range(PICTURES):
    p = shots / f"repo{i}__light.png"
    # Distinct pixels per picture, so "one part per distinct picture" cannot pass by two pictures
    # happening to compress to the same bytes.
    PILImage.new("RGB", (300, 150), (7 + 60 * i, 40, 90)).save(p)
    SRC.append(p)

renders = []


def render(src: Path) -> io.BytesIO:
    """Stands in for `07_build.pad_shot`: one JPEG canvas per call, and it records the calls."""
    renders.append(str(src))
    buf = io.BytesIO()
    with PILImage.open(src) as im:
        im.convert("RGB").save(buf, "JPEG", quality=88)
    buf.seek(0)
    return buf


def anchor(row: int) -> TwoCellAnchor:
    return TwoCellAnchor(editAs="twoCell",
                         _from=AnchorMarker(col=2, colOff=0, row=row - 1, rowOff=0),
                         to=AnchorMarker(col=3, colOff=0, row=row, rowOff=0))


def fresh() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def names(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return z.namelist()


# ---- 1. what openpyxl does on its own: a part per placement, not per picture
plain = fresh()
for s in range(SHEETS):
    ws = plain.create_sheet(f"S{s}")
    for i, src in enumerate(SRC):
        img = XLImage(render(src))
        img.anchor = anchor(2 + i)
        ws.add_image(img)
p_plain = TMP / "plain.xlsx"
plain.save(p_plain)
plain_names = names(p_plain)
plain_media = [n for n in plain_names if n.startswith("xl/media/")]
eq("openpyxl writes one media part per placement", len(plain_media), PLACEMENTS)
eq("...so it renders once per placement too", len(renders), PLACEMENTS)

# ---- 2. the pool: one part per distinct picture, and nothing else moves
renders.clear()
pool = media.Pool(render)
pooled = fresh()
for s in range(SHEETS):
    ws = pooled.create_sheet(f"S{s}")
    for i, src in enumerate(SRC):
        pool.place(ws, src, anchor(2 + i))
p_pooled = TMP / "pooled.xlsx"
report = media.save(pooled, p_pooled, pool)
pooled_names = names(p_pooled)
pooled_media = [n for n in pooled_names if n.startswith("xl/media/")]

eq("the pool reports every placement", report.placements, PLACEMENTS)
eq("the pool reports one image per distinct source", report.parts, PICTURES)
eq("the archive agrees about the part count", len(pooled_media), PICTURES)
eq("each picture is rendered once, not once per cell", len(renders), PICTURES)
eq("the report's entry count is the archive's", report.entries, len(pooled_names))
eq("headroom is the rest of the ceiling", report.headroom, media.ENTRY_LIMIT - report.entries)

# The whole ceiling argument is that media are the only part of the package that grows, so the two
# files differing *only* in their media is the claim worth asserting -- not a formula for the rest.
eq("deduplication changes nothing but the media parts",
   sorted(n for n in pooled_names if not n.startswith("xl/media/")),
   sorted(n for n in plain_names if not n.startswith("xl/media/")))
eq("and the entry counts differ by exactly the duplicates removed",
   len(plain_names) - len(pooled_names), PLACEMENTS - PICTURES)

# Named for the pool slot, so the name does not depend on which sheet was written first.
eq("parts are named for the image, not the placement", sorted(pooled_media),
   [f"xl/media/{media.PART_STEM}{i}.jpeg" for i in range(1, PICTURES + 1)])

# Sharing works through one `Default` for the extension. A per-image `Override` would defeat it, and
# would be the first thing to break if openpyxl changed how it declares media.
with zipfile.ZipFile(p_pooled) as z:
    ctypes = z.read("[Content_Types].xml").decode("utf-8")
eq("jpeg is declared once, by extension", ctypes.count('Extension="jpeg"'), 1)
true("no part-specific Override for any picture", "/xl/media/" not in ctypes, ctypes[:400])

# ---- 3. the package still reads back, with every picture still on every sheet
back = load_workbook(p_pooled)
eq("every sheet survives", back.sheetnames, [f"S{s}" for s in range(SHEETS)])
eq("every placement survives", [len(back[s]._images) for s in back.sheetnames], [PICTURES] * SHEETS)
eq("twoCell anchoring survives -- it is what hides a picture with its filtered row",
   sorted({type(i.anchor).__name__ for s in back.sheetnames for i in back[s]._images}),
   ["TwoCellAnchor"])
seen = {bytes(i._data()) for i in back["S0"]._images}
eq("and the four pictures read back as four different pictures", len(seen), PICTURES)

# ---- 4. every drawing relationship resolves to a part that is in the file
# The failure this rules out is the interesting one: sharing a part means a relationship written while
# sheet 1 was serialised has to still name a part that exists after sheet 5 has been written.
present = set(pooled_names)
dangling, embeds = [], 0
with zipfile.ZipFile(p_pooled) as z:
    for drawing in sorted(n for n in pooled_names if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n)):
        rels_part = f"xl/drawings/_rels/{Path(drawing).name}.rels"
        rels = z.read(rels_part).decode("utf-8") if rels_part in present else ""
        # `Id` and `Target` appear in either order depending on openpyxl's serialiser, so pair them per
        # element rather than by two separate scans.
        target = {}
        for el in re.findall(r"<Relationship [^>]*/?>", rels):
            rid = re.search(r'Id="([^"]+)"', el)
            tgt = re.search(r'Target="([^"]+)"', el)
            if rid and tgt:
                target[rid.group(1)] = tgt.group(1)
        used = re.findall(r'r:embed="([^"]+)"', z.read(drawing).decode("utf-8"))
        embeds += len(used)
        for rid in used:
            t = target.get(rid)
            if t is None:
                dangling.append((drawing, rid, "no such relationship"))
                continue
            # Targets are written absolute (`/xl/media/...`); a relative one would resolve against the
            # drawing's own directory, so both spellings are handled rather than assumed.
            part = t[1:] if t.startswith("/") else \
                os.path.normpath(os.path.join("xl/drawings", t)).replace("\\", "/")
            if part not in present:
                dangling.append((drawing, rid, part))
true("no picture points at a part that is not in the package", not dangling, str(dangling[:4]))
eq("the drawings between them hold every placement", embeds, PLACEMENTS)

# ---- 5. one file spelled two ways is still one part
# `shots_all.json` stores basenames the build joins onto `cache/shots`; the original pipeline's
# `shots.json` stores whole paths it wrote itself. The same picture can arrive either way.
renders.clear()
pool2 = media.Pool(render)
wb2 = fresh()
ws = wb2.create_sheet("S")
pool2.place(ws, SRC[0], anchor(2))
pool2.place(ws, SRC[0].parent / "." / SRC[0].name, anchor(3))
if os.name == "nt":
    # Windows only: `normcase` folds case there and nowhere else, because nowhere else may it.
    pool2.place(ws, Path(str(SRC[0]).upper()), anchor(4))
eq("a normalised path is the same image", pool2.parts, 1)
eq("...and it is rendered once", len(renders), 1)
true("...however many times it is placed", pool2.placements >= 2, str(pool2.placements))

# ---- 6. the two guards
# The ceiling guard. Checked by moving the ceiling rather than by writing 65,536 pictures: the
# comparison is the thing under test, and one is not slow.
real_limit = media.ENTRY_LIMIT
media.ENTRY_LIMIT = report.entries - 1
try:
    p_over = TMP / "over.xlsx"
    pool3 = media.Pool(render)
    wb3 = fresh()
    for s in range(SHEETS):
        ws = wb3.create_sheet(f"S{s}")
        for i, src in enumerate(SRC):
            pool3.place(ws, src, anchor(2 + i))
    msg = raises("the ceiling guard fires", lambda: media.save(wb3, p_over, pool3))
    true("...naming the ceiling and the counts", str(media.ENTRY_LIMIT) in msg.replace(",", ""), msg)
    true("...and leaves no half-written workbook behind", not p_over.exists())
finally:
    media.ENTRY_LIMIT = real_limit

# The postcondition. `save` compares the pool against the archive because the deduplication rides on
# openpyxl internals (`ExcelWriter._images`, `Image._data`) and CI installs openpyxl unpinned -- an
# upstream rename would put the per-placement writer back with nothing to announce it. A plain
# `XLImage` alongside the pooled ones stands in for that: one more part in the file than in the pool.
p_extra = TMP / "extra.xlsx"
pool4 = media.Pool(render)
wb4 = fresh()
ws = wb4.create_sheet("S")
pool4.place(ws, SRC[0], anchor(2))
stray = XLImage(render(SRC[1]))
stray.anchor = anchor(3)
ws.add_image(stray)
msg = raises("a part the pool does not know about is caught", lambda: media.save(wb4, p_extra, pool4))
true("...and says deduplication is not happening", "not happening" in msg, msg)
true("...without leaving the workbook", not p_extra.exists())

# Without a pool `save` is still `wb.save` plus the ceiling check, which is what `07_build` standing on
# its own wants: one sheet, no other sheet to share with, no postcondition to check.
p_solo = TMP / "solo.xlsx"
wb5 = fresh()
ws = wb5.create_sheet("S")
solo = XLImage(render(SRC[2]))
solo.anchor = anchor(2)
ws.add_image(solo)
solo_report = media.save(wb5, p_solo)
eq("no pool, no deduplication, no complaint", solo_report.parts, 1)
eq("a plain image keeps openpyxl's own part name, which is why PART_STEM is not `image`",
   [n for n in names(p_solo) if n.startswith("xl/media/")], ["xl/media/image1.jpeg"])

# ---- 7. the projection: is the 65,535 ceiling actually reachable, and where?
#
# Published or measured, not assumed:
LISTS_TODAY = 11                 # README, and `docs/data.json`'s `listed_by` values
SHEETED_ENTRIES = 2_394          # the eleven per-list sheets: 426+264+255+219+215+194+194+193+165+153+116
REPOS_TODAY = 1_294              # rows in `docs/data.json` -- the By Category sheet is one row each
LEADERBOARD = 250                # `16_build_all.main` slices the board to 250
PLATFORM_ROWS_TODAY = 3_527      # Windows 1,105 + macOS 1,092 + Linux 1,120 + Docker 210, from data.json
MEASURED_PLACEMENTS = 7_388      # `18_slicers.inject`'s docstring: the one count from a real build
# Non-media parts, by the rules asserted in section 2 above: 8 package-level, one per worksheet, three
# more (sheet rels, drawing, drawing rels) for each of the 18 picture sheets and 2 chart sheets, four
# charts, and the By Category table. Two orders of magnitude below the media count, which is the point.
FIXED = 8 + 21 + 3 * 20 + 4 + 1

# Sheet rows per repo on the four platform sheets. The one number the answer turns on, so it is measured
# from the published data rather than assumed: a repo qualifies for 2.73 of the four sheets on average,
# not 4, because Docker reaches 16% of them.
PLATFORM_RATIO = PLATFORM_ROWS_TODAY / REPOS_TODAY
true("a repo lands on 2.7 of the four platform sheets, not 4",
     2.7 < PLATFORM_RATIO < 2.75, f"{PLATFORM_RATIO:.4f}")


def placements(lists: int, repos: int, ratio: float = PLATFORM_RATIO) -> int:
    """Pictures placed in one theme's workbook.

    Only the cross-list sheets grow with the atlas. Sources gains a row per list; the eleven per-list
    sheets are the eleven lists that have sheets, and ingesting a twelfth list does not add a twelfth
    sheet -- new lists reach readers through By Category, the Leaderboard and the platform sheets.
    """
    return lists + SHEETED_ENTRIES + LEADERBOARD + repos + round(ratio * repos)


def entries(lists: int, repos: int, ratio: float = PLATFORM_RATIO, shared: bool = False) -> int:
    """ZIP entries in that workbook. Shared media means one part per picture, i.e. per repo."""
    return FIXED + (repos if shared else placements(lists, repos, ratio))


# The model has one measurement to answer to, and it is within 1.2% of it. The shortfall is rows whose
# screenshot file is missing and platform rows excluded as subpaths, both of which make it an
# over-estimate -- the safe direction for a ceiling.
modelled = placements(LISTS_TODAY, REPOS_TODAY)
true("the model reproduces the one real measurement",
     0 <= modelled - MEASURED_PLACEMENTS <= MEASURED_PLACEMENTS * 0.015,
     f"modelled {modelled:,} vs measured {MEASURED_PLACEMENTS:,}")

# Where PR #2 lands: 39 lists, 13,640 listings, 7,980 distinct repos (its own summary table).
PR2_LISTS, PR2_REPOS = 39, 7_980
pr2 = entries(PR2_LISTS, PR2_REPOS)
true("the 39-list merge is nowhere near the ceiling",
     pr2 < media.ENTRY_LIMIT // 2, f"{pr2:,} entries, {pr2 / media.ENTRY_LIMIT:.1%}")

# And where JFH-214 lands. 7,835 *rows*, not repos, and these lists overlap heavily -- 775 of today's
# repos are already on more than one list -- so treating every row as a new repo is the worst case.
TB_LISTS, TB_REPOS = 40, PR2_REPOS + 7_835
tb = entries(TB_LISTS, TB_REPOS)
true("even assuming no overlap at all, TensorBlock stays under the ceiling",
     tb < media.ENTRY_LIMIT, f"{tb:,} entries, {tb / media.ENTRY_LIMIT:.1%}")
true("...but only just: under 10% of the ceiling left",
     0 < media.ENTRY_LIMIT - tb < media.ENTRY_LIMIT * 0.1, f"{media.ENTRY_LIMIT - tb:,} spare")

# Assume instead that every repo qualifies for all four platform sheets and the ticket's conclusion
# follows, at around the 118% it quotes -- the exact figure depends on how many of the 7,835 rows are
# repos the atlas does not already have. That single assumption is the whole distance between "past the
# ceiling" and "94% of it", and it is the one the published data contradicts: Docker reaches 16% of
# repos, and a list of MCP servers pushes that ratio down rather than up.
all_four = entries(TB_LISTS, TB_REPOS, ratio=4.0)
true("assuming all four platform sheets per repo is what puts it past the ceiling",
     1.10 < all_four / media.ENTRY_LIMIT < 1.30, f"{all_four:,}, {all_four / media.ENTRY_LIMIT:.1%}")
true("...and under that assumption the ceiling arrives before TensorBlock does",
     entries(TB_LISTS, 12_600, ratio=4.0) > media.ENTRY_LIMIT > entries(TB_LISTS, 12_400, ratio=4.0),
     f"{entries(TB_LISTS, 12_600, ratio=4.0):,}")


def crossing(shared: bool) -> int:
    """Distinct repos at which one more repo would take the workbook past the ceiling."""
    lo, hi = 1, 1_000_000
    while lo < hi:
        mid = (lo + hi) // 2
        if entries(40, mid, shared=shared) > media.ENTRY_LIMIT:
            hi = mid
        else:
            lo = mid + 1
    return lo - 1


# So the ceiling is real and it is close -- roughly one more large list past JFH-214, not several -- and
# sharing the media moves it out by a factor of the placement ratio, to about four times the size the
# atlas would be with TensorBlock in it.
per_placement, per_picture = crossing(False), crossing(True)
true("per placement, the ceiling arrives at about 16,800 repos",
     16_500 < per_placement < 17_200, f"{per_placement:,}")
true("shared, it arrives at about 65,000",
     60_000 < per_picture < 66_000, f"{per_picture:,}")
true("...which is past TensorBlock by 4x, not by 6%",
     per_picture > 4 * TB_REPOS > per_placement, f"{per_picture:,} vs {TB_REPOS:,}")

# ---- 8. the build's own choke point, on a synthetic row set
# `16_build_all.place_shots` is where six of the seven sheet families put their pictures, so this is the
# call the projection above is really about. Importing the module is safe without the cache -- it only
# reads `cache/` inside `main()` -- and `SHOTS` is the one thing that has to be redirected.
spec = importlib.util.spec_from_file_location(
    "b16", Path(__file__).resolve().parent.parent / "scripts" / "16_build_all.py")
b16 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b16)
b16.SHOTS = shots

# Three repos with a screenshot on file, one whose file was never captured. The missing one is not a
# hypothetical: the modelled placement count overshoots the real workbook by 1.2% precisely because rows
# like it are skipped.
rows = [{"shot_key": f"repo{i}"} for i in range(PICTURES)] + [{"shot_key": "nofile"}]
shot_map = {f"repo{i}": {"light": SRC[i].name, "dark": SRC[i].name} for i in range(PICTURES)}
shot_map["nofile"] = {"light": "never-captured.png", "dark": ""}

renders.clear()
pool5 = media.Pool(render)
wb6 = fresh()
# Every sheet family that shows a repo shows the same repos, which is the whole reason for the pool.
for sheet in ("Leaderboard", "By Category", "Windows", "macOS", "Linux", "Docker"):
    b16.place_shots(wb6.create_sheet(sheet), rows, shot_map, "light", pool5, 3, 4)
p_sheets = TMP / "sheets.xlsx"
sheets_report = media.save(wb6, p_sheets, pool5)

eq("a row whose screenshot was never captured places nothing", sheets_report.placements, 6 * PICTURES)
eq("six sheets of the same repos still embed each picture once", sheets_report.parts, PICTURES)
eq("...and still render each picture once", len(renders), PICTURES)
# The arithmetic of section 7, on a workbook small enough to count by hand: 8 package parts, one per
# sheet, three more per sheet carrying pictures, and one part per distinct picture.
eq("the entry count is the fixed cost plus one part per picture",
   sheets_report.entries, 8 + 6 + 3 * 6 + PICTURES)

print(f"\n  today ({LISTS_TODAY} lists, {REPOS_TODAY:,} repos): "
      f"{entries(LISTS_TODAY, REPOS_TODAY):,} entries per placement, "
      f"{entries(LISTS_TODAY, REPOS_TODAY, shared=True):,} shared")
print(f"  PR #2 ({PR2_LISTS} lists, {PR2_REPOS:,} repos): {pr2:,} "
      f"({pr2 / media.ENTRY_LIMIT:.0%} of the ceiling), "
      f"{entries(PR2_LISTS, PR2_REPOS, shared=True):,} shared")
print(f"  + TensorBlock ({TB_REPOS:,} repos, worst case): {tb:,} "
      f"({tb / media.ENTRY_LIMIT:.0%}), {entries(TB_LISTS, TB_REPOS, shared=True):,} shared")
print(f"  ceiling at {per_placement:,} repos per placement, {per_picture:,} shared")

print(f"\n{ok} passed, {bad} failed")
sys.exit(1 if bad else 0)
