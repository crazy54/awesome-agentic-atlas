"""The two workbook covers carry the same small Atlas mark and mascot as the site.

This builds only the cover in a disposable workbook. The real build needs a complete crawler cache,
which is intentionally not part of a normal checkout, so the test keeps the branding regression local
to the generator rather than requiring a release-sized workbook.
"""
from __future__ import annotations

import importlib.util
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "16_build_all.py"

sys.path.insert(0, str(ROOT / "scripts"))
import media  # noqa: E402  -- after the path insert, which is what makes it importable

spec = importlib.util.spec_from_file_location("b16_branding", SCRIPT)
b16 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b16)

checks = 0


def ok(condition: bool, label: str) -> None:
    global checks
    checks += 1
    if not condition:
        raise AssertionError(label)


STATS = {
    "lists": 1, "items": 1, "repos": 1, "stars": 1, "shots": 0, "multi": 0,
    "pool": 1, "platforms": {"Windows": 1, "macOS": 1, "Linux": 1, "Docker": 0},
    "site_rows": 0, "sub_rows": 0, "unavailable": 0,
}
SOURCES = [{
    "sheet": "Example List", "items": 1, "sum_stars": 1, "nwo": "example/list",
    "blurb": "A disposable source row used only to check the workbook cover.",
}]

ok(b16.BRAND_MARK.is_file(), "the square Atlas logo source asset exists")
ok(b16.BRAND_MASCOT.is_file(), "the Atlas Byte source asset exists")

with tempfile.TemporaryDirectory(prefix="atlas-workbook-branding-") as tmp:
    tmp = Path(tmp)
    for theme in ("dark", "light"):
        wb = Workbook()
        wb.remove(wb.active)
        ws = b16.build_cover(wb, b16.THEMES[theme], STATS, SOURCES)
        path = tmp / f"{theme}.xlsx"
        wb.save(path)

        reopened = load_workbook(path, read_only=False)
        cover = reopened["Start Here"]
        ok(len(cover._images) == 2, f"{theme} cover writes exactly the logo and mascot")
        # openpyxl reports a reopened image's natural pixel size. Its OneCellAnchor extent is the
        # displayed size, in EMUs (9,525 per pixel), so assert that rather than the source bitmap.
        dims = sorted((image.anchor.ext.width // 9525, image.anchor.ext.height // 9525)
                      for image in cover._images)
        ok(dims == [(32, 32), (86, 90)], f"{theme} cover keeps compact artwork bounds")
        anchors = sorted((image.anchor._from.col, image.anchor._from.row) for image in cover._images)
        ok(anchors == [(11, 0), (12, 0)], f"{theme} cover keeps artwork in the upper-right margin")
        ok(cover["M5"].value == "ATLAS BYTE", f"{theme} cover labels the mascot")
        ok(cover["K6"].value and "Switch to the" in cover["K6"].value,
           f"{theme} cover retains the sibling-theme control")
        with zipfile.ZipFile(path) as archive:
            drawings = [name for name in archive.namelist() if name.startswith("xl/drawings/drawing")]
            embedded = [name for name in archive.namelist() if name.startswith("xl/media/")]
        ok(len(drawings) == 1, f"{theme} cover has one drawing part")
        ok(len(embedded) == 2, f"{theme} cover embeds two brand images")

        # The release build does not save with plain openpyxl; it saves through `media.save`, whose
        # postcondition compares the screenshot pool against the package. That is the half of JFH-292
        # this file owns, and until it the two had never been in one process -- `media_test.py` builds
        # workbooks with no cover, and everything above here saves with `wb.save`. So save a cover the
        # way a weekly does, and let them meet in two seconds here rather than 74 minutes into a run.
        through_pool = tmp / f"{theme}-pooled.xlsx"
        pooled_wb = Workbook()
        pooled_wb.remove(pooled_wb.active)
        pooled_cover = b16.build_cover(pooled_wb, b16.THEMES[theme], STATS, SOURCES)
        # Empty, and its renderer is never called, because the cover's artwork is deliberately not
        # pooled: `Pool` renders through `pad_shot`, which pads a *screenshot* onto a theme-coloured
        # canvas, and it sizes each placement from the bytes -- while `cover_image` bounds the mark to
        # 32px and the mascot to 86x90 explicitly. There is also nothing to share them with: two
        # placements, one cover, one each.
        pool = media.Pool(lambda src: (_ for _ in ()).throw(AssertionError("cover must not be pooled")))
        ok(not any(isinstance(image, media.SharedImage) for image in pooled_cover._images),
           f"{theme} cover artwork is placed outside the pool, which is what lets it be bounded")
        report = media.save(pooled_wb, through_pool, pool)
        ok(report.parts == 2 and report.placements == 2,
           f"{theme} cover survives the pooling writer instead of reading as failed deduplication")

        # The fallback operates on release assets when a local checkout lacks the crawler cache.
        # Start with a bare cover so this also proves the package surgery adds a drawing relationship
        # without needing an existing release-sized workbook in the test fixture.
        fallback = tmp / f"{theme}-fallback.xlsx"
        bare = Workbook()
        bare.active.title = "Start Here"
        bare.save(fallback)
        b16.inject_existing_cover_brand(fallback)
        b16.verify_existing_cover_brand(fallback)
        fallback_cover = load_workbook(fallback, read_only=False)["Start Here"]
        ok(len(fallback_cover._images) == 2, f"{theme} fallback writes both cover images")
        with zipfile.ZipFile(fallback) as archive:
            ok(archive.testzip() is None, f"{theme} fallback leaves a readable XLSX archive")
            ok(sum(name.startswith("xl/media/") for name in archive.namelist()) == 2,
               f"{theme} fallback adds exactly two media parts")

print(f"{checks} passed, 0 failed")
